"""Documents on request: "make me a PDF / Word file / spreadsheet about ..." is a deliverable, not a
coding project.

The local model writes the content as Markdown in the requested language; Coding Brain converts it
to HTML itself (everything escaped: the model's text can never become markup or script) and prints
PDFs with the system's Chromium/Edge through Playwright, offline and with JavaScript disabled.
Word files are written as plain WordprocessingML (no macros), spreadsheets with openpyxl (values only,
never formulas), CSV and text as UTF-8. Images and slide decks are not produced locally.
Files go to Documents\\Coding Brain (CODINGBRAIN_DOCUMENTS overrides) and are never overwritten.
"""
from __future__ import annotations

import html
import os
import re
import time
from pathlib import Path

FORMATS = {"pdf": "pdf", "docx": "docx", "xlsx": "xlsx", "csv": "csv", "txt": "txt", "html": "html", "md": "md"}
FORMAT_WORDS = [(r"\b(word|docx|\.doc|ms\s*word)\b", "docx"), (r"\b(excel|xlsx|spreadsheet|workbook)\b", "xlsx"),
                (r"\bcsv\b", "csv"), (r"\b(txt|plain\s+text|text\s+file|e-?mail)\b", "txt"),
                (r"\b(markdown|\.md|md)\b", "md"), (r"\b(html|web\s*page)\b", "html"), (r"\bpdf\b", "pdf")]
TABULAR = {"xlsx", "csv"}
UNSUPPORTED = re.compile(r"(?i)\b(image|picture|photo|logo|drawing|illustration|png|jpe?g|svg|icon|"
                         r"power\s*point|pptx|slides?|slide\s+deck|video|audio|mp3|mp4)\b")
LANGUAGES = {"french": "French", "français": "French", "francais": "French", "english": "English",
             "anglais": "English", "spanish": "Spanish", "español": "Spanish", "espagnol": "Spanish",
             "german": "German", "deutsch": "German", "allemand": "German", "portuguese": "Portuguese",
             "italian": "Italian", "italien": "Italian", "arabic": "Arabic", "chinese": "Chinese"}
LENGTHS = [(r"\b(one|1|a single|single)\s+page\b|\bshort\b|\bbrief\b", 900), (r"\blong\b|\bdetailed\b|\b(5|6|7|8|9|10)\s+pages\b", 2600)]
CSS = """
@page { size: A4; margin: 22mm 20mm; }
body { font: 11pt/1.55 "Segoe UI", Calibri, Arial, sans-serif; color: #1d2324; }
h1 { font-size: 22pt; margin: 0 0 6mm; color: #0f3d33; } h2 { font-size: 15pt; margin: 7mm 0 2mm; color: #155e4c; }
h3 { font-size: 12.5pt; margin: 5mm 0 1.5mm; } p { margin: 0 0 3mm; } ul, ol { margin: 0 0 3mm 6mm; }
table { border-collapse: collapse; width: 100%; margin: 3mm 0 4mm; font-size: 10pt; }
th, td { border: 1px solid #c9d6d2; padding: 2mm 2.5mm; text-align: left; vertical-align: top; }
th { background: #e8f2ef; } code { font-family: Consolas, monospace; background: #f1f4f3; padding: 0 1mm; }
pre { background: #f1f4f3; padding: 3mm; white-space: pre-wrap; font: 9.5pt Consolas, monospace; }
hr { border: 0; border-top: 1px solid #c9d6d2; margin: 5mm 0; } .meta { color: #6b7b78; font-size: 9pt; }
"""

TABLE_PROMPT = """You prepare spreadsheet data for the user, in {language}. Output Markdown only: for each
sheet a "## " heading with the sheet name, then one Markdown table with a header row. Plain values only:
no formulas, no HTML, no commentary. Use realistic example values only when the user asked for an
example or template, and say "example" in the sheet name then. About {words} cells in total at most."""

WRITER_PROMPT = """You write documents for the user. Write the complete document the user asked for, in {language}.
Use Markdown only: one "# " title, "## " sections, short paragraphs, "- " lists and Markdown tables where
useful. No HTML, no code fences unless the document is about code, no notes to the user, no preamble:
start with the title. Be accurate; when a fact is uncertain, write it in general terms instead of
inventing specifics such as names, numbers or quotes. Aim for about {words} words."""


class DocumentError(RuntimeError):
    pass


def detect_format(text: str) -> str:
    lowered = text.lower()
    for pattern, name in FORMAT_WORDS:
        if re.search(pattern, lowered):
            return name
    return "pdf"


def detect_language(text: str) -> str | None:
    lowered = text.lower()
    for word, name in LANGUAGES.items():
        if re.search(rf"\b{re.escape(word)}\b", lowered):
            return name
    return None


def target_words(text: str) -> int:
    for pattern, tokens in LENGTHS:
        if re.search(pattern, text.lower()):
            return tokens * 3 // 4
    return 1500 * 3 // 4


def output_folder() -> Path:
    base = os.environ.get("CODINGBRAIN_DOCUMENTS")
    return Path(base) if base else Path.home() / "Documents" / "Coding Brain"


def _inline(text: str) -> str:
    text = html.escape(text, quote=True)
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"(?<![*\w])\*([^*\n]+)\*(?!\w)", r"<em>\1</em>", text)
    return text


def markdown_to_html(markdown: str, title: str = "Document") -> str:
    """A small, safe Markdown subset: headings, paragraphs, lists, tables, code blocks, rules."""
    out, paragraph, lines, index = [], [], markdown.replace("\r\n", "\n").split("\n"), 0

    def flush():
        if paragraph:
            out.append("<p>" + " ".join(_inline(line) for line in paragraph) + "</p>")
            paragraph.clear()
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            block, index = [], index + 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            out.append("<pre>" + html.escape("\n".join(block)) + "</pre>")
        elif re.match(r"^#{1,6}\s", stripped):
            flush()
            level = min(len(stripped) - len(stripped.lstrip("#")), 6)
            out.append(f"<h{level}>{_inline(stripped[level:].strip())}</h{level}>")
        elif re.match(r"^(-{3,}|\*{3,})$", stripped):
            flush()
            out.append("<hr>")
        elif re.match(r"^([-*+]|\d+[.)])\s+", stripped):
            flush()
            ordered = bool(re.match(r"^\d", stripped))
            items = []
            while index < len(lines) and re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[index]):
                items.append("<li>" + _inline(re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", lines[index])) + "</li>")
                index += 1
            out.append(("<ol>" if ordered else "<ul>") + "".join(items) + ("</ol>" if ordered else "</ul>"))
            continue
        elif stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[index + 1]):
            flush()
            cells = lambda row: [cell.strip() for cell in row.strip().strip("|").split("|")]  # noqa: E731
            head = "".join(f"<th>{_inline(cell)}</th>" for cell in cells(stripped))
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append("<tr>" + "".join(f"<td>{_inline(cell)}</td>" for cell in cells(lines[index])) + "</tr>")
                index += 1
            out.append(f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(rows)}</tbody></table>")
            continue
        elif not stripped:
            flush()
        else:
            paragraph.append(stripped)
        index += 1
    flush()
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<style>{CSS}</style></head><body>{''.join(out)}</body></html>")


def title_of(markdown: str, fallback: str) -> str:
    match = re.search(r"^#\s+(.+)$", markdown, re.M)
    return (match.group(1) if match else fallback).strip()[:120]


def _slug(text: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE).strip().lower()
    return re.sub(r"[\s_-]+", "-", slug)[:60].strip("-") or "document"


def unique_path(folder: Path, title: str, extension: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    base = f"{_slug(title)}-{time.strftime('%Y%m%d-%H%M')}"
    candidate = folder / f"{base}.{extension}"
    number = 2
    while candidate.exists():
        candidate = folder / f"{base}-{number}.{extension}"
        number += 1
    return candidate


def print_pdf(page_html: str, target: Path, launch_options: dict) -> None:
    from playwright.sync_api import sync_playwright
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(**launch_options)
        try:
            context = browser.new_context(java_script_enabled=False, offline=True, accept_downloads=False)
            context.route("**/*", lambda route: route.abort() if not route.request.url.startswith(("data:", "about:"))
                          else route.continue_())
            page = context.new_page()
            page.set_content(page_html, wait_until="load")
            page.pdf(path=str(target), format="A4", print_background=True, prefer_css_page_size=True)
        finally:
            browser.close()


def _cells(row: str) -> list[str]:
    return [cell.strip() for cell in row.strip().strip("|").split("|")]


def tables_of(markdown: str) -> list[tuple[str, list[list[str]]]]:
    """(sheet name, rows) for each Markdown table, named after the heading above it."""
    sheets, name, lines, index = [], None, markdown.splitlines(), 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith("#"):
            name = line.lstrip("#").strip()
        elif line.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[index + 1]):
            rows, index = [_cells(line)], index + 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_cells(lines[index]))
                index += 1
            sheets.append((name or f"Sheet{len(sheets) + 1}", rows))
            continue
        index += 1
    return sheets


def _value(text: str):
    """Numbers become numbers; nothing ever becomes a formula."""
    cleaned = text.replace(",", "").replace("$", "").replace("€", "").strip()
    if re.fullmatch(r"-?\d+", cleaned):
        return int(cleaned)
    if re.fullmatch(r"-?\d+\.\d+", cleaned):
        return float(cleaned)
    return text.lstrip("=+@") if text[:1] in "=+@" else text  # never a formula or DDE


def write_xlsx(markdown: str, target: Path, title: str) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    sheets = tables_of(markdown) or [(title, [[line] for line in markdown.splitlines() if line.strip()])]
    book = Workbook()
    book.remove(book.active)
    used = set()
    for name, rows in sheets:
        safe = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Sheet"
        while safe in used:
            safe = (safe[:28] + f" {len(used)}")[:31]
        used.add(safe)
        sheet = book.create_sheet(safe)
        for row in rows:
            sheet.append([_value(cell) for cell in row])
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="E8F2EF")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(10, width + 2), 60)
    book.save(target)
    return len(sheets)


def write_csv(markdown: str, target: Path) -> None:
    import csv
    sheets = tables_of(markdown)
    rows = sheets[0][1] if sheets else [[line] for line in markdown.splitlines() if line.strip()]
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            writer.writerow([cell.lstrip("=+-@") if cell[:1] in "=+@" else cell for cell in row])


def _docx_run(text: str, bold=False, italic=False, size=None, mono=False) -> str:
    props = ("<w:b/>" if bold else "") + ("<w:i/>" if italic else "") + \
            (f'<w:sz w:val="{size}"/>' if size else "") + \
            ('<w:rFonts w:ascii="Consolas" w:hAnsi="Consolas"/>' if mono else "")
    return f'<w:r>{f"<w:rPr>{props}</w:rPr>" if props else ""}<w:t xml:space="preserve">{html.escape(text, quote=False)}</w:t></w:r>'


def _docx_inline(text: str, **base) -> str:
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![*\w])\*[^*\n]+\*(?!\w))", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append(_docx_run(part[2:-2], bold=True, **base))
        elif part.startswith("`"):
            out.append(_docx_run(part[1:-1], mono=True, **base))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(_docx_run(part[1:-1], italic=True, **base))
        else:
            out.append(_docx_run(part, **base))
    return "".join(out)


def write_docx(markdown: str, target: Path, title: str) -> None:
    """A minimal, valid .docx (WordprocessingML): headings, paragraphs, lists and tables. No macros."""
    import zipfile
    body, lines, index = [], markdown.replace("\r\n", "\n").split("\n"), 0
    para = lambda inner, spacing=120: f'<w:p><w:pPr><w:spacing w:after="{spacing}"/></w:pPr>{inner}</w:p>'  # noqa: E731
    while index < len(lines):
        line = lines[index].strip()
        if re.match(r"^#{1,6}\s", line):
            level = len(line) - len(line.lstrip("#"))
            size = {1: 40, 2: 30, 3: 26}.get(level, 24)
            body.append(para(_docx_inline(line[level:].strip(), size=size).replace("<w:rPr>", "<w:rPr><w:b/>")
                             .replace("<w:r><w:t", "<w:r><w:rPr><w:b/></w:rPr><w:t"), 160))
        elif re.match(r"^([-*+]|\d+[.)])\s+", line):
            marker = "•" if line[0] in "-*+" else re.match(r"^\d+", line).group(0) + "."
            body.append(f'<w:p><w:pPr><w:ind w:left="360" w:hanging="240"/></w:pPr>{_docx_run(marker + " ")}'
                        f'{_docx_inline(re.sub(r"^([-*+]|\d+[.)])\s+", "", line))}</w:p>')
        elif line.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[index + 1]):
            rows, index = [_cells(line)], index + 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_cells(lines[index]))
                index += 1
            border = '<w:tblBorders>' + "".join(f'<w:{side} w:val="single" w:sz="4" w:color="C9D6D2"/>'
                                                for side in ("top", "left", "bottom", "right", "insideH", "insideV")) + '</w:tblBorders>'
            table = [f'<w:tbl><w:tblPr><w:tblW w:w="0" w:type="auto"/>{border}</w:tblPr>']
            for number, row in enumerate(rows):
                table.append("<w:tr>" + "".join(f'<w:tc><w:p>{_docx_inline(cell, bold=number == 0)}</w:p></w:tc>' for cell in row) + "</w:tr>")
            body.append("".join(table) + "</w:tbl>" + para(""))
            continue
        elif line and not re.match(r"^(-{3,}|```)", line):
            body.append(para(_docx_inline(line)))
        index += 1
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                + "".join(body) + '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/>'
                '<w:pgMar w:top="1247" w:right="1134" w:bottom="1247" w:left="1134"/></w:sectPr></w:body></w:document>')
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml",
                         '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                         '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                         '<Default Extension="xml" ContentType="application/xml"/>'
                         '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                         '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/></Types>')
        package.writestr("_rels/.rels",
                         '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                         '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/></Relationships>')
        package.writestr("docProps/core.xml",
                         '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
                         'xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>' + html.escape(title) + '</dc:title><dc:creator>Coding Brain</dc:creator></cp:coreProperties>')
        package.writestr("word/document.xml", document)


def page_count(path: Path) -> int | None:
    try:
        import pypdfium2
        document = pypdfium2.PdfDocument(str(path))
        try:
            return len(document)
        finally:
            document.close()
    except Exception:
        return None


def create(model, request: str, details: str = "", fmt: str | None = None, language: str | None = None,
           launch_options: dict | None = None, step=lambda text: None, on_token=None) -> dict:
    """Write the document with the local model and save it. Returns {path, format, title, pages, bytes, words}."""
    wanted = f"{request}\n{details}".strip()
    if UNSUPPORTED.search(request) and not re.search(r"(?i)\b(pdf|word|docx|excel|xlsx|csv|text|markdown|html)\b", wanted):
        raise DocumentError("images, slide decks, audio and video cannot be produced locally; ask for a PDF or Word "
                            "handout instead, or describe the picture and I can write the text that goes with it")
    fmt = FORMATS.get((fmt or detect_format(wanted)).lower(), "pdf")
    language = language or detect_language(wanted) or "the language the user wrote in"
    words = target_words(wanted)
    step(f"Writing the document with {getattr(model, 'model', None) or 'the local model'}")
    prompt = (TABLE_PROMPT if fmt in TABULAR else WRITER_PROMPT).format(language=language, words=words)
    messages = [{"role": "system", "content": prompt},
                {"role": "user", "content": f"Request: {request}" + (f"\nDetails: {details}" if details else "")}]
    max_tokens = min(int(words * 1.6) + 200, 3600)
    markdown = (model.chat(messages, max_tokens=max_tokens, on_token=on_token) if on_token
                else model.chat(messages, max_tokens=max_tokens)).strip()
    markdown = re.sub(r"^```(?:markdown|md)?\s*\n|\n```\s*$", "", markdown)
    if len(markdown) < 40:
        raise DocumentError("the model did not write the document; try again or give more details")
    title = title_of(markdown, request)
    folder = output_folder()
    target = unique_path(folder, title, fmt)
    sheets = None
    if fmt == "md":
        target.write_text(markdown + "\n", encoding="utf-8")
    elif fmt == "txt":
        plain = re.sub(r"^#{1,6}\s*", "", markdown, flags=re.M)
        target.write_text(re.sub(r"\*\*([^*]+)\*\*", r"\1", plain) + "\n", encoding="utf-8")
    elif fmt == "docx":
        step("Creating the Word document")
        write_docx(markdown, target, title)
    elif fmt == "xlsx":
        step("Creating the spreadsheet")
        sheets = write_xlsx(markdown, target, title)
    elif fmt == "csv":
        write_csv(markdown, target)
    else:
        page = markdown_to_html(markdown, title)
        if fmt == "html":
            target.write_text(page, encoding="utf-8")
        else:
            step("Creating the PDF")
            print_pdf(page, target, launch_options or {"headless": True})
    return {"path": str(target), "format": fmt, "title": title, "pages": page_count(target) if fmt == "pdf" else None,
            "sheets": sheets, "bytes": target.stat().st_size, "words": len(markdown.split())}
